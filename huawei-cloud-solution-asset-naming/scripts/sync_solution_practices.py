#!/usr/bin/env python3
"""Sync Huawei Cloud solution-practice cards into the skill workbook."""

import argparse
import html
import json
import re
from copy import copy
from html.parser import HTMLParser
from pathlib import Path
from urllib.error import HTTPError
from urllib.parse import urlparse, urlunparse
from urllib.request import Request, urlopen

from openpyxl import Workbook, load_workbook


SITES = {
    "中国站": "https://www.huaweicloud.com/solution/reference-architecture.html",
    "国际站": "https://www.huaweicloud.com/intl/zh-cn/solution/reference-architecture.html",
}
HEADERS = ["序号", "主标签", "一级场景", "二级场景", "解决方案名称", "方案名或实践名（英文）", "ID编码", "方案站点", "中国站区域", "国际站区域", "中国站链接", "国际站链接", "简述"]
LINK_HEADERS = ("中国站链接", "国际站链接")
TITLE_ALIASES = [{"litellm统一的ai管理网关", "litellm统一的ai网关"}]


class CardsParser(HTMLParser):
    def __init__(self):
        super().__init__()
        self.cards = None

    def handle_starttag(self, _tag, attrs):
        if self.cards is not None:
            return
        value = dict(attrs).get("data-cards")
        if value:
            self.cards = json.loads(html.unescape(value))


def fetch_cards(site, url):
    request = Request(url, headers={"User-Agent": "Mozilla/5.0 solution-practice-sync/1.0"})
    with urlopen(request, timeout=30) as response:
        body = response.read().decode("utf-8")
    parser = CardsParser()
    parser.feed(body)
    if parser.cards is None:
        raise RuntimeError(f"{site}页面未找到 data-cards，网页结构可能已变化：{url}")
    cards = []
    for item in parser.cards:
        card = item.get("cardItem", {})
        if card.get("caption") and card.get("href"):
            target = official_url(card["href"])
            if not target:
                raise RuntimeError(f"{site}方案链接格式异常，已停止同步：{clean_text(card['caption'])} {card['href']}")
            card["href"] = target
            card["site"] = site
            cards.append(card)
    return cards


def clean_text(value):
    value = re.sub(r"<[^>]+>", "", html.unescape(value or ""))
    return re.sub(r"\s+", " ", value).strip()


def normalized_title(value):
    return re.sub(r"[\s,，:：;；—–-]+", "", value or "").lower()


def slug(value):
    return Path(urlparse(value or "").path).stem.lower()


def official_url(value):
    value = html.unescape(str(value or "")).strip()
    parsed = urlparse(value)
    if (parsed.scheme != "https" or parsed.netloc.lower() != "www.huaweicloud.com"
            or not parsed.path.endswith(".html") or parsed.username or parsed.password):
        return ""
    return urlunparse(parsed._replace(fragment=""))


def check_url(url):
    request = Request(url, headers={"User-Agent": "Mozilla/5.0 solution-practice-sync/1.0"}, method="HEAD")
    try:
        with urlopen(request, timeout=15) as response:
            return response.status, response.geturl()
    except HTTPError as error:
        if error.code not in (405, 501):
            raise
        request = Request(url, headers={"User-Agent": "Mozilla/5.0 solution-practice-sync/1.0", "Range": "bytes=0-0"})
        with urlopen(request, timeout=15) as response:
            return response.status, response.geturl()


def naming_notes(name, main_label="应用方案"):
    notes = []
    if main_label == "应用方案" and any(word in name for word in ("快速部署", "一键部署", "免费试用")):
        notes.append("应用方案名称含“快速部署/一键部署/免费试用”等禁用宣传词")
    if re.search(r"[,，:：;；—–]", name):
        notes.append("名称使用逗号、冒号或破折号，不符合单空格分隔规则")
    return "；".join(notes)


def same_solution(row, card):
    card_urls = [card.get("href", ""), card.get("官网链接", "")]
    card_urls.extend(card.get(header, "") for header in LINK_HEADERS)
    card_slugs = {slug(url) for url in card_urls if url}
    row_urls = [row.get("href", ""), row.get("官网链接", "")]
    row_urls.extend(row.get("_links", {}).values())
    row_urls.extend(row.get(header, "") for header in LINK_HEADERS)
    if card_slugs & {slug(url) for url in row_urls if url}:
        return True
    left = normalized_title(row.get("解决方案名称") or row.get("caption"))
    right = normalized_title(card.get("caption") or card.get("解决方案名称"))
    if left == right:
        return True
    return any({left, right} <= aliases for aliases in TITLE_ALIASES)


def interleave_site_cards(site_cards):
    # shortcut: one workbook row order cannot mirror both catalogs; add per-site ranks if both must be exact.
    china = site_cards.get("中国站", [])
    intl = site_cards.get("国际站", [])
    before, after, tail = {}, {}, []

    def china_index(card):
        return next((i for i, item in enumerate(china) if same_solution(item, card)), None)

    for index, card in enumerate(intl):
        if china_index(card) is not None:
            continue
        following = next((china_index(item) for item in intl[index + 1:] if china_index(item) is not None), None)
        if following is not None:
            before.setdefault(following, []).append(card)
            continue
        previous = next((china_index(item) for item in reversed(intl[:index]) if china_index(item) is not None), None)
        if previous is None:
            tail.append(card)
        else:
            after.setdefault(previous, []).append(card)

    ordered = []
    for index, card in enumerate(china):
        ordered.extend(before.get(index, []))
        ordered.append(card)
        ordered.extend(after.get(index, []))
    ordered.extend(tail)
    return ordered


def merge_cards(cards):
    merged = []
    for card in cards:
        match = next((item for item in merged if same_solution(item, card)), None)
        if match:
            match["sites"].add(card["site"])
            match[f"{card['site']}链接"] = card["href"]
        else:
            merged.append({
                "中国站链接": card["href"] if card["site"] == "中国站" else "",
                "国际站链接": card["href"] if card["site"] == "国际站" else "",
                "解决方案名称": clean_text(card["caption"]),
                "一级场景": clean_text(card.get("label", "")).split(",")[0],
                "简述": clean_text(card.get("description", "")),
                "sites": {card["site"]},
            })
    return merged


def audit_order(site_cards, merged):
    reports = []
    for site, cards in site_cards.items():
        positions = []
        missing = []
        for card in cards:
            position = next((i for i, row in enumerate(merged) if same_solution(row, card)), None)
            if position is None:
                missing.append(clean_text(card["caption"]))
            else:
                positions.append((position, clean_text(card["caption"])))
        inversions = [(left, right) for left, right in zip(positions, positions[1:]) if left[0] >= right[0]]
        reports.append(f"{site}目录：官网 {len(cards)} 条，匹配 {len(positions)} 条，顺序相邻冲突 {len(inversions)} 处")
        reports.extend(f"  {site}未匹配：{name}" for name in missing)
        reports.extend(f"  {site}顺序冲突：{left[1]} → {right[1]}" for left, right in inversions[:5])

    china = site_cards.get("中国站", [])
    intl = site_cards.get("国际站", [])
    for index, card in enumerate(intl):
        if any(same_solution(item, card) for item in china):
            continue
        previous = next((item for item in reversed(intl[:index]) if any(same_solution(c, item) for c in china)), None)
        following = next((item for item in intl[index + 1:] if any(same_solution(c, item) for c in china)), None)
        if previous and following:
            previous_index = next(i for i, item in enumerate(china) if same_solution(item, previous))
            following_index = next(i for i, item in enumerate(china) if same_solution(item, following))
            if previous_index > following_index:
                reports.append(
                    f"  两站锚点冲突：{clean_text(card['caption'])} 在国际站位于"
                    f"{clean_text(previous['caption'])}与{clean_text(following['caption'])}之间，"
                    "但中国站顺序相反；已按规则放在后一个国际站共有方案之前。"
                )
    return reports


def site_text(sites):
    return "中国站、国际站" if sites == {"中国站", "国际站"} else next(iter(sites), "")


def copy_row_style(sheet, source_row, target_row):
    for column in range(1, sheet.max_column + 1):
        source, target = sheet.cell(source_row, column), sheet.cell(target_row, column)
        if source.has_style:
            target._style = copy(source._style)
        target.number_format = source.number_format
        target.alignment = copy(source.alignment)
    sheet.row_dimensions[target_row].height = sheet.row_dimensions[source_row].height


def load_sheet(path):
    if path.exists():
        workbook = load_workbook(path)
        return workbook, workbook.active
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "解决方案实践"
    sheet.append(HEADERS)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:M1"
    return workbook, sheet


def workbook_urls(path):
    if not path.exists():
        return [], []
    sheet = load_workbook(path, read_only=False, data_only=False).active
    headers = [cell.value for cell in sheet[1]]
    urls = []
    invalid = []
    for row in sheet.iter_rows(min_row=2):
        for header in (*LINK_HEADERS, "官网链接"):
            if header not in headers:
                continue
            cell = row[headers.index(header)]
            raw = cell.hyperlink.target if cell.hyperlink else cell.value
            url = official_url(raw)
            if url:
                urls.append(url)
            elif raw:
                invalid.append((row[headers.index("解决方案名称")].value if "解决方案名称" in headers else "", header, raw))
    return urls, invalid


def sync(path, cards, records=()):
    existed = path.exists()
    workbook, sheet = load_sheet(path)
    freeze_panes = sheet.freeze_panes
    headers = [cell.value for cell in sheet[1]]
    for header in HEADERS:
        if header not in headers:
            sheet.cell(1, len(headers) + 1, header)
            headers.append(header)
    indexes = {header: headers.index(header) + 1 for header in headers}
    rows = []
    invalid_links = []
    for cells in sheet.iter_rows(min_row=2):
        row = {header: cells[column - 1].value or "" for header, column in indexes.items()}
        row["_links"] = {
            header: official_url(
                sheet.cell(cells[0].row, indexes[header]).hyperlink.target
                if header in indexes and sheet.cell(cells[0].row, indexes[header]).hyperlink
                else row.get(header, "")
            )
            for header in LINK_HEADERS
        }
        for header in LINK_HEADERS:
            cell = sheet.cell(cells[0].row, indexes[header]) if header in indexes else None
            raw = cell.hyperlink.target if cell and cell.hyperlink else row.get(header, "")
            if raw and not row["_links"][header]:
                invalid_links.append((row.get("解决方案名称", ""), header, raw))
        if "官网链接" in indexes:
            legacy_cell = sheet.cell(cells[0].row, indexes["官网链接"])
            legacy = official_url(legacy_cell.hyperlink.target if legacy_cell.hyperlink else row.get("官网链接", ""))
            if legacy and not any(row["_links"].values()):
                site = "国际站" if "/intl/" in urlparse(legacy).path else "中国站"
                row["_links"][f"{site}链接"] = legacy
                row["_links"]["官网链接"] = legacy
        row["_sites"] = {site for header, site in zip(LINK_HEADERS, ("中国站", "国际站")) if row["_links"][header]}
        rows.append(row)

    added = 0
    added_names = []
    ordered_rows = []
    matched_rows = set()
    for card in merge_cards(cards):
        match = next((row for row in rows if same_solution(row, card)), None)
        if match:
            match["_sites"] = set(card["sites"])
            match["_links"] = {header: card.get(header, "") for header in LINK_HEADERS}
        else:
            match = {
                "主标签": "应用方案",
                "一级场景": card["一级场景"],
                "二级场景": "",
                "解决方案名称": card["解决方案名称"],
                "方案名或实践名（英文）": "",
                "ID编码": slug(card.get("中国站链接") or card.get("国际站链接")),
                "中国站区域": "",
                "国际站区域": "",
                "中国站链接": card.get("中国站链接", ""),
                "国际站链接": card.get("国际站链接", ""),
                "简述": card["简述"],
                "其他标签": "",
                "_sites": card["sites"],
                "_links": {header: card.get(header, "") for header in LINK_HEADERS},
            }
            added += 1
            added_names.append(match["解决方案名称"])
        ordered_rows.append(match)
        matched_rows.add(id(match))
    stale_rows = [row for row in rows if id(row) not in matched_rows]
    rows = ordered_rows + stale_rows

    for record in records:
        match = next((row for row in rows if same_solution(row, record)), None)
        if match:
            match.update({key: value for key, value in record.items() if value and key in HEADERS})
            match.setdefault("_links", {}).update({header: official_url(record.get(header)) for header in LINK_HEADERS if record.get(header)})
        else:
            match = {header: record.get(header, "") for header in HEADERS}
            match["_links"] = {header: official_url(record.get(header)) for header in LINK_HEADERS}
            rows.insert(0, match)
            added += 1
            added_names.append(match.get("解决方案名称", ""))
        for header in LINK_HEADERS:
            raw = record.get(header, "")
            if raw and not official_url(raw):
                invalid_links.append((record.get("解决方案名称", ""), header, raw))
        match["_sites"] = {site for header, site in zip(LINK_HEADERS, ("中国站", "国际站")) if match.get("_links", {}).get(header)}

    for number, row in enumerate(rows, 1):
        excel_row = number + 1
        if excel_row > sheet.max_row and sheet.max_row > 1:
            copy_row_style(sheet, sheet.max_row, excel_row)
        row["序号"] = number
        row["方案站点"] = site_text(row.get("_sites", set()))
        notes = [value for value in (row.get("备注", ""), naming_notes(row.get("解决方案名称", ""), row.get("主标签", ""))) if value]
        if not row.get("方案名或实践名（英文）") or not row.get("二级场景"):
            if "备注" in indexes:
                notes.append("现网新增，待补充英文名称和二级场景")
        if "备注" in indexes:
            row["备注"] = "；".join(dict.fromkeys(notes))
        for header, column in indexes.items():
            cell = sheet.cell(excel_row, column)
            if header in LINK_HEADERS or header == "官网链接":
                target = row.get("_links", {}).get(header) or row.get(header, "")
                if target and not str(target).startswith(("https://", "http://")):
                    target = ""
                cell.value = "查看方案" if target else ""
                cell.hyperlink = target or None
            else:
                cell.value = row.get(header, "")

    sheet.auto_filter.ref = f"A1:{sheet.cell(1, sheet.max_column).column_letter}{sheet.max_row}"
    sheet.auto_filter.filterColumn = []
    for row_number in range(2, sheet.max_row + 1):
        sheet.row_dimensions[row_number].hidden = False
    sheet.freeze_panes = freeze_panes if existed else "A2"
    sheet.sheet_view.selection[0].activeCell = "A1"
    sheet.sheet_view.selection[0].sqref = "A1"
    region_gaps = [
        (row.get("解决方案名称", ""), site)
        for row in rows
        for link_header, region_header, site in (
            ("中国站链接", "中国站区域", "中国站"),
            ("国际站链接", "国际站区域", "国际站"),
        )
        if row.get("_links", {}).get(link_header) and not row.get(region_header)
    ]
    workbook.save(path)
    return len(rows), added, added_names, [row.get("解决方案名称", "") for row in stale_rows], invalid_links, region_gaps


def main():
    default = Path(__file__).resolve().parents[1] / "解决方案实践.xlsx"
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", nargs="?", type=Path, default=default)
    parser.add_argument("--record-json", action="append", default=[], help="要新增或更新的一条完整表格记录（JSON 对象）")
    parser.add_argument("--check-links", action="store_true", help="同步前检查官网目录和工作簿中所有官网链接的可访问性")
    args = parser.parse_args()
    site_cards = {site: list(reversed(fetch_cards(site, url))) for site, url in SITES.items()}
    cards = interleave_site_cards(site_cards)
    merged = merge_cards(cards)
    for line in audit_order(site_cards, merged):
        print(line)
    if args.check_links:
        current_urls, invalid_existing = workbook_urls(args.workbook)
        if invalid_existing:
            print(f"工作簿中有 {len(invalid_existing)} 条格式无效或缺少可识别超链接的官网链接；未修改工作簿：")
            for name, header, value in invalid_existing:
                print(f"  {name}（{header}）：{value}")
            return 2
        urls = list(dict.fromkeys(
            [card["href"] for site_list in site_cards.values() for card in site_list]
            + current_urls
        ))
        failures = []
        for url in urls:
            try:
                status, final_url = check_url(url)
                if not 200 <= status < 400:
                    failures.append((url, f"HTTP {status}"))
            except Exception as error:
                failures.append((url, str(error)))
        if failures:
            print(f"官网链接检查失败 {len(failures)}/{len(urls)} 条；未修改工作簿：")
            for url, reason in failures:
                print(f"  {url} — {reason}")
            return 2
        print(f"官网链接检查通过：{len(urls)} 条。")
    records = [json.loads(value) for value in args.record_json]
    total, added, added_names, stale, invalid_links, region_gaps = sync(args.workbook, cards, records)
    print(f"已同步 {total} 条方案实践，新增 {added} 条：{args.workbook}")
    if added_names:
        print(f"新增：{', '.join(added_names)}")
    if stale:
        print(f"保留 {len(stale)} 条不在两站目录中的既有行：{', '.join(stale)}")
    if invalid_links:
        print(f"已清除 {len(invalid_links)} 条格式无效的官网链接：")
        for name, header, value in invalid_links:
            print(f"  {name}（{header}）：{value}")
    if added:
        print("新增方案的区域、英文名称和二级场景需按详情页及命名规则核实补齐；目录页不提供这些字段。")
    if region_gaps:
        print(f"有 {len(region_gaps)} 个已发布站点链接缺少对应区域，需按该站详情页核实：")
        for name, site in region_gaps[:10]:
            print(f"  {name}（{site}）")
        if len(region_gaps) > 10:
            print(f"  另有 {len(region_gaps) - 10} 项")
    return 0


if __name__ == "__main__":
    main()
